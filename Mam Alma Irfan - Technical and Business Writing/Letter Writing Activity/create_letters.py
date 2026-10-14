from docx import Document
from docx.shared import Inches, Pt


OUTPUT_DIR = r"e:\Study Material\5th Semester All Data\Mam Alma Irfan - Technical and Business Writing\Letter Writing Activity"


def set_document_style(document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def add_paragraph(document, text="", bold=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return paragraph


def add_letter(document, sender_address, date, recipient_address, subject,
               salutation, body, closing, sender_name, sender_role,
               enclosure=None):
    add_paragraph(document, sender_address)
    add_paragraph(document, date)
    add_paragraph(document, recipient_address)
    add_paragraph(document, f"Subject: {subject}", bold=True)
    add_paragraph(document, salutation)

    for paragraph_text in body:
        add_paragraph(document, paragraph_text)

    add_paragraph(document, closing)
    add_paragraph(document, "")
    add_paragraph(document, "")
    add_paragraph(document, "")
    add_paragraph(document, sender_name)
    add_paragraph(document, sender_role)
    if enclosure:
        add_paragraph(document, f"Enclosure: {enclosure}")


combined = Document()
set_document_style(combined)
add_letter(
    combined,
    "Department of Software Engineering\nInstitute of Management Sciences\nPeshawar, PAKISTAN",
    "October 14, 2026",
    "Mr. Salman Khan\nManager\nTechZone Electronics\nUniversity Road\nPeshawar, PAKISTAN",
    "Complaint about defective wireless headphones",
    "Dear Mr. Khan:",
    [
        "I am writing to complain about the wireless headphones that I purchased from TechZone Electronics on October 10, 2026. The product is SoundWave Pro Wireless Headphones, and the invoice number is TZ-1048.",
        "Unfortunately, the headphones do not work properly. The left speaker produces no sound, and the battery stops working after approximately thirty minutes of use. I followed the charging and setup instructions, but the problem has continued.",
        "As the product is still under warranty, I kindly request a replacement pair or a full refund. I have enclosed a copy of the invoice for your reference. Please let me know how I should return the defective product.",
        "I hope you will resolve this matter promptly. I would appreciate your response within seven working days. Thank you for your attention to this complaint.",
    ],
    "Sincerely,",
    "Abdul Samad",
    "BSSE, 5th Semester\nInstitute of Management Sciences",
    "Copy of invoice",
)
combined.add_page_break()
add_letter(
    combined,
    "TechZone Electronics\nUniversity Road\nPeshawar, PAKISTAN",
    "October 17, 2026",
    "Mr. Abdul Samad\nBSSE, 5th Semester\nInstitute of Management Sciences\nPeshawar, PAKISTAN",
    "Response to your complaint about wireless headphones",
    "Dear Mr. Samad:",
    [
        "Thank you for writing to us about the SoundWave Pro Wireless Headphones purchased under invoice number TZ-1048. We sincerely apologize for the inconvenience caused by the faulty speaker and battery.",
        "We have reviewed your complaint and confirm that the product is covered by our warranty. We are pleased to offer you a replacement pair at no additional cost. You may return the defective headphones to our store with a copy of the invoice.",
        "Once we receive the product, the replacement will be ready within three working days. If you prefer a refund instead, please inform us when you return the headphones, and we will process it according to our refund policy.",
        "We value your business and appreciate you bringing this matter to our attention. Please accept our apologies, and contact us if you need any further assistance.",
    ],
    "Sincerely,",
    "Salman Khan",
    "Manager\nTechZone Electronics",
)
combined.save(OUTPUT_DIR + r"\Business Letter Activity - Complaint and Reply.docx")

print("Created the combined Word document.")