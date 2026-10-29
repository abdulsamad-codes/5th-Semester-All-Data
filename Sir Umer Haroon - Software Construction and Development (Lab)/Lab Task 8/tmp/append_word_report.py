from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import json
import shutil
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
target = root / 'Abdul Samad - Lab Task 8 - 5th Semester Group A.docx'
backup = root / 'tmp' / 'original_word.docx'
if not backup.exists():
    shutil.copy2(target, backup)
results = json.loads((root / 'test_results.json').read_text())
doc = Document()

def p(text='', bold=False, size=10, before=0, after=5, keep=False):
    para = doc.add_paragraph()
    f = para.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing = 1.05
    f.keep_with_next = keep
    run = para.add_run(text)
    run.font.name, run.font.size = 'Verdana', Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    return para

def h(text):
    return p(text, True, 12, 6, 6, True)

def page():
    para = p(after=0)
    para.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell,width in zip(t.rows[0].cells,widths): cell.width = Inches(width)
    for col,width in zip(t.columns,widths): col.width = Inches(width)
    for row in [headers]+rows:
        cells = t.rows[0].cells if row is headers else t.add_row().cells
        for cell,text,width in zip(cells,row,widths):
            cell.width = Inches(width)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.line_spacing = 1.0
            r=para.add_run(str(text)); r.font.name='Verdana';r.font.size=Pt(8)
            r.bold = row is headers; r.font.color.rgb=RGBColor(0,0,0)
            tcPr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
            tcPr.append(mar)
        trPr=cells[0]._tc.getparent().get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'000000');borders.append(el)
    t._tbl.tblPr.append(borders)
    p(after=1)

page()
h('Task 1: The Crash Gallery')
p('I tested the original functions with bad inputs, recorded the errors, and then added checks to handle those inputs safely. The original and patched functions are included in crash_gallery.py.')
h('Original results')
first=[]
for name in ['divide','get_first','get_age','parse_number','read_file','average']:
    first += [r for r in results['task_1'] if r['function']==name and r['bad']][:3]
table(['Bad input','What happened','Error type'],[[r['call'],r['original'],r['error_type']] for r in first],[2.2,3.0,1.35])
p('The test block catches each original exception so the remaining tests can still run. Temporary file paths are shown as missing.txt and <directory>.',size=9)
page()
h('Task 1: Patched results')
p('I ran the same bad inputs again after adding the checks. All of them returned clear messages instead of crashing.')
table(['Input','Result after patching'],[[r['call'],r['patched'].strip("'")] for r in first],[2.5,4.05])
h('Discussion')
p('1. Hardest function: read_file was the hardest because a valid filename can still refer to a missing file, a directory or a file that cannot be opened. I used path checks and a small I/O exception handler for read-time failures.')
p('2. Common bad input: Wrong types were the most common, such as None instead of a list or text instead of a number.')
p('3. Pattern in the fixes: I checked the type first, then checked the value or contents before doing the operation.')
p('4. More than one check: Yes. For example, divide needs numeric inputs and a nonzero divisor, while average needs a nonempty list containing numbers.')
page()
h('Task 2: The Silent Bug Hunt')
p('I tested normal values and edge cases to compare the expected and original results. The corrected functions and comments explaining each fix are in bug_hunt.py.')
p("For search tests: students = [{'name': 'Ali', 'age': 20}, {'name': 'Sara', 'age': 19}]. Marks use a 0 to 100 scale, and a tax rate of 0.1 means 10%.",size=9)
second=[]
for name in ['calculate_discount','calculate_average','find_student','is_adult','apply_tax','count_passed']:
    second += [r for r in results['task_2'] if r['function']==name][:3]
def shorten(value):
    return value.replace("'Error: ",'').rstrip("'") if value.startswith("'Error: ") else value
table(['Test input','Expected result','Actual original result','Bug?'],[[r['call'],shorten(r['expected']),r['original'],'Yes' if r['bug'] else 'No'] for r in second],[2.05,2.0,2.1,.4])
p('The fixed result matched the expected result in every row. For invalid input, the fixed functions return an Error message; is_adult returns False and an unsuccessful student search returns None.',size=9)
page()
h('Task 2: Bug report')
for name,text in [
('calculate_discount','It should apply a discount between 0 and 100%. The original returned a negative price for a 150% discount. I added checks for the price and discount range.'),
('calculate_average','It should calculate the mean of valid marks. The original crashed on an empty list and accepted marks outside the allowed range. I checked the list and marks before calculating the average.'),
('find_student','It should return the matching student or None. The original returned the first student when the name was missing. I removed that fallback and checked that each record has a name.'),
('is_adult','It should count age 18 as adult. The original returned False for 18 because it used > 18. I changed the comparison to >= 18 and rejected invalid ages.'),
('apply_tax','It should add tax to the amount. The original accepted a negative tax rate and reduced the amount. I added checks to reject negative amounts and tax rates.'),
('count_passed','It should count marks equal to or above the passing mark. The original did not count marks exactly equal to 50. I used >= instead of > and validated the marks and threshold.')]:
    p(name,True,10,before=3,after=3,keep=True);p(text)
h('Discussion')
p('1. Hardest bug to find: find_student was the hardest because it returned a student that looked valid even when the requested name did not exist.')
p('2. How edge cases helped: Testing age 18 and marks of exactly 50 exposed the comparison errors. Empty lists and negative values also showed problems that normal inputs did not reveal.')
p('3. Why silent bugs are dangerous: A wrong result can look correct and be used in later work. A crash is easier to notice because it stops the program and shows an error.')
h('Final result')
p('I ran both Python files after making the fixes. Task 1 passed 26 checks and Task 2 passed 31 checks, including the cases shown above. All 57 checks passed.')

with ZipFile(backup) as z:
    entries={i.filename:(i,z.read(i.filename)) for i in z.infolist()}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
xml=etree.fromstring(entries['word/document.xml'][1])
body=xml.find('w:body',ns)
# Keep the user's cover and closing paragraphs, inserting only before the
# closing separator and The End.
paragraphs=body.findall('w:p',ns)
anchor=paragraphs[-2]
index=body.index(anchor)
for el in list(doc.element.body):
    if el.tag!=qn('w:sectPr'):
        body.insert(index,deepcopy(el));index+=1
new_xml=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
output=root/'tmp'/'completed_word.docx'
with ZipFile(output,'w',ZIP_DEFLATED) as z:
    for name,(info,data) in entries.items():
        z.writestr(info,new_xml if name=='word/document.xml' else data)
shutil.copy2(output,target)
with ZipFile(target) as z:
    assert all(z.read(n)==data for n,(_,data) in entries.items() if n!='word/document.xml')
print('Updated',target.name)
print('Preserved every original package part except the document body.')
print('Added 36 test cases, patched results, six short bug reports and all discussion answers.')
