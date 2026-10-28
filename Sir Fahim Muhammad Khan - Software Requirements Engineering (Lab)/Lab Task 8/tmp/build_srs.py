from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output'
OUT.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.8)
sec.left_margin = sec.right_margin = Inches(.85)
sec.header_distance = sec.footer_distance = Inches(.35)
sec.different_first_page_header_footer = True

styles = doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Caption']:
    st = styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'Times New Roman')
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Times New Roman')
styles['Normal'].font.size = Pt(11.5)
styles['Normal'].paragraph_format.line_spacing = 1.1
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.widow_control = True
for name,size in [('Title',25),('Heading 1',17),('Heading 2',13),('Heading 3',11.5),('Heading 4',11.5)]:
    styles[name].font.size=Pt(size)
    styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(12 if name!='Title' else 0)
    styles[name].paragraph_format.space_after=Pt(6)
    styles[name].paragraph_format.keep_with_next=True
styles['Subtitle'].font.size=Pt(15)
styles['Subtitle'].font.italic=False
front=styles.add_style('Contents Heading',1)
front.base_style=styles['Normal']
front.font.name='Times New Roman';front.font.size=Pt(17);front.font.bold=True
front.paragraph_format.space_after=Pt(10)
styles['Caption'].font.size=Pt(10)
styles['Caption'].font.italic=True

def p(text='',bold=False):
    x=doc.add_paragraph()
    r=x.add_run(text);r.bold=bold
    return x

headings=[]
def h(text,level=1,new=False):
    x=doc.add_paragraph(text,'Contents Heading' if text=='Table of Contents' else f'Heading {level}')
    if new:x.paragraph_format.page_break_before=True
    headings.append(text)
    return x

def label(name,text):
    x=doc.add_paragraph()
    x.add_run(name+' ').bold=True
    x.add_run(text)
    return x

def bullets(items):
    for item in items:
        x=doc.add_paragraph(style='List Bullet');x.add_run(item)
        x.paragraph_format.space_after=Pt(3)

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,val in enumerate(headers):t.rows[0].cells[i].text=val
    for row in rows:
        cells=t.add_row().cells
        for c,val in zip(cells,row):c.text=str(val)
    for j,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
        if j==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for i,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[i])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6' if j==0 else 'FFFFFF');pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.space_before=Pt(2)
                para.paragraph_format.line_spacing=1.02
                if j==0:para.paragraph_format.keep_with_next=True
                if i==0 and len(str(c.text))<22:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(10.5);run.font.bold=(j==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def field(paragraph,instruction):
    r=OxmlElement('w:r');begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r.append(begin);paragraph._p.append(r)
    r=OxmlElement('w:r');it=OxmlElement('w:instrText');it.set(qn('xml:space'),'preserve');it.text=instruction;r.append(it);paragraph._p.append(r)
    r=OxmlElement('w:r');sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r.append(sep);paragraph._p.append(r)
    r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text='';r.append(tx);paragraph._p.append(r)
    r=OxmlElement('w:r');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r.append(end);paragraph._p.append(r)

header=sec.header.paragraphs[0]
header.text='Enterprise RAG Assistant                                      Software Requirements Specification'
for r in header.runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Version 1.0  |  Page ');field(footer,' PAGE ');footer.add_run(' of ');field(footer,' NUMPAGES ')
for r in footer.runs:r.font.size=Pt(9)

# Cover
p('Institute of Management Sciences, Peshawar',True).alignment=WD_ALIGN_PARAGRAPH.CENTER
p('BS Software Engineering').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Software Requirements Engineering Lab').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('')
x=doc.add_paragraph('Enterprise RAG Assistant','Title');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
x=doc.add_paragraph('Software Requirements Specification','Subtitle');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Version 1.0').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Complete planned version 1').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('')
for line in ['Prepared by','Abdul Samad','Student ID 243203624','Submitted to','Sir Fahim Muhammad Khan','27 October 2026']:
    x=p(line,bold=line in ['Abdul Samad','Sir Fahim Muhammad Khan']);x.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('')
p('Status: Prepared for review and approval').alignment=WD_ALIGN_PARAGRAPH.CENTER

h('Document Control',1,True)
p('This specification defines the intended first release of Enterprise RAG Assistant. It provides a common basis for implementation, testing, academic review, and acceptance. It does not certify that the planned functions have already been implemented or that the quality targets have been achieved.')
table(['Item','Value'],[
('Document identifier','ERA-SRS-001'),('Product','Enterprise RAG Assistant'),('Version and date','1.0 | 27 October 2026'),('Author and owner','Abdul Samad | Student ID 243203624'),('Institution','Institute of Management Sciences, Peshawar'),('Programme and course','BS Software Engineering | Software Requirements Engineering Lab'),('Reviewer','Sir Fahim Muhammad Khan'),('Approval status','Pending the records in Section 5'),('Scope baseline','Complete planned version 1; one organisation and four simulated employee roles')],[1.55,5.25])
h('Revision History',2)
table(['Version','Date','Description','Author'],[('1.0','27 October 2026','Initial formal SRS assembled from the project requirements v0.2, design v0.3, data-source records, and the supplied course guidance.','Abdul Samad')],[.65,1.05,3.8,1.3])
label('Interpretation.','“Shall” identifies a required behaviour. “Should” identifies a desired behaviour. “May” identifies an optional capability. Items labelled provisional retain their original target but require the measurement decisions in Section 6.3 before formal acceptance.')
label('Requirement control.','The existing FR-01 to FR-13 and NFR-01 to NFR-07 identifiers are retained. Supporting rules use BR, IR, DR, and DC identifiers. These rules make the existing scope explicit; they do not create additional product releases or user roles.')
label('Priority.','High means required for the version 1 baseline. Medium means a provisional quality target awaiting calibration. These priorities are proposed for review; scope changes follow Section 4.')

h('Table of Contents',1,True)
toc=p();field(toc,' TOC \\o "1-3" \\h \\z \\u ')

h('1 Introduction',1,True)
h('1.1 Purpose',2)
p('Enterprise RAG Assistant helps employees find answers in an approved collection of company documents. This Software Requirements Specification defines the functions, access rules, external interfaces, data obligations, and quality targets for its first release. It also states how each requirement will be verified.')
p('The intended readers are the project owner, the course instructor, developers, testers, and anyone reviewing the portfolio demonstration. The owner uses this document to control scope; developers use Section 3 to implement the required behaviour; reviewers use Section 6 to check coverage and identify unresolved decisions.')
h('1.2 Scope',2)
p('The product accepts employee questions in English and produces answers grounded in documents the employee is permitted to access. Every substantive answer includes a source citation. When the available evidence cannot support an answer, the assistant returns a generic not-found response. Its purpose is to reduce manual document searching while preserving access boundaries and making answers traceable.')
p('Version 1 includes offline ingestion of Markdown and text-based PDF documents, department tagging, authenticated question answering, a REST API, a browser-based chat interface, query metrics, a repeatable evaluation process, and a deployed demonstration. The demonstration represents a single fictional company, Acme Corp, using selected public GitLab Handbook pages.')
p('Voice interaction, multilingual support, a native mobile application, model training or fine-tuning, editing source documents through chat, live synchronisation with company systems, enterprise SSO, and multiple-organisation hosting are outside the first release. OCR for scanned PDFs is a possible later extension and is not an acceptance obligation for version 1.')
h('1.3 Definitions Acronyms and Abbreviations',2)
table(['Term','Meaning in this specification'],[
('SRS','Software Requirements Specification; the controlled statement of required product behaviour.'),
('RAG','Retrieval-augmented generation; answering a question using document passages retrieved for that question.'),
('LLM','Large language model used to produce an answer from the supplied evidence.'),
('RBAC','Role-based access control; permission determined from an authenticated user role.'),
('Document and chunk','A source file and a smaller passage derived from that file. A chunk retains its document identity and department.'),
('Citation','A reference to an authorised source, including its document title and section or PDF page.'),
('Not found','The same generic outcome for insufficient evidence, restricted material, or non-existent material; it does not reveal whether a restricted document exists.'),
('Faithfulness','The degree to which an answer is supported by its supplied source passages, as scored by the agreed evaluation method.'),
('API and bearer token','A software interface and the signed access token included with authenticated requests.'),
('Checksum','A SHA-256 fingerprint used to identify a source document revision.'),
('BM25 and reranker','Keyword-based retrieval scoring and a later step that reorders candidate passages for relevance.'),
('RRF','Reciprocal rank fusion; a method of merging ranked retrieval results.'),
('FR and NFR','Functional requirement and non-functional requirement.'),
('BR IR DR DC','Business rule, interface rule, data rule, and design or operating constraint.'),
('TBD','To be decided and recorded before the affected acceptance criterion is approved.')],[1.35,5.45])
h('1.4 References',2)
p('The following sources establish the project scope and document structure. Local filenames identify the supplied copies. The course guidance references ISO/IEC/IEEE 29148; the template states that it was adapted from IEEE Std 830-1993. This document applies the supplied guidance and does not claim independently audited conformity to either standard.')
table(['Ref','Document and source','Version or date'],[
('R1','Software Requirement Specification (SRS), supplied course handout, SRS.pdf.','Four-page supplied copy; issue date not stated.'),
('R2','Software Requirements Specification (SRS) Template, CS3911, supplied annotated PDF.','30-page copy with header date 08/15/22.'),
('R3','Requirements Document: Enterprise RAG Assistant, Abdul Samad, docs/01-requirements.md.','Version 0.2.'),
('R4','Design Document: Enterprise RAG Assistant, Abdul Samad, docs/02-design.md.','Version 0.3.'),
('R5','Data Sources and dataset manifest, docs/data-sources.md and docs/data-manifest.csv.','Pinned handbook revision; downloaded 11 October 2026.'),
('R6','Project README, AGENTS.md, source modules, and tests in the local project repository.','Local review on 27 October 2026.'),
('R7','GitLab Handbook, GitLab: handbook.gitlab.com; source repository gitlab.com/gitlab-com/content-sites/handbook.','Commit 2c77a1f5b8c8a80cb7b5151ff11cfad84f98bfbb.'),
('R8','GitLab Handbook licence attribution recorded in the project README and data-source record.','CC BY-SA 4.0; creativecommons.org/licenses/by-sa/4.0/.')],[.45,4.4,1.95])
p('Project repository: https://github.com/abdulsamad-codes/enterprise-rag-assistant. References R3 to R6 are available in the local project at F:\\Industry\\enterprise-rag-assistant. The course PDFs are retained with this SRS in the lab task folder.')
h('1.5 Overview',2)
p('Section 2 explains the product boundary, users, environment, and release scope. Section 3 specifies observable behaviour and the conditions used to verify it. Sections 4 and 5 control changes and approval. Section 6 provides the acceptance plan, traceability, open decisions, and supporting operational records. Technology selections remain in the separate design document except where an established operating constraint affects acceptance.')

h('2 The Overall Description',1,True)
h('2.1 Product Perspective',2)
p('The assistant is a standalone document-question-answering application for one organisation. It supplements the source handbook rather than replacing its policies or editing workflow. It does not act on payroll, approve leave, change financial records, or perform the business transactions described in the documents. Users remain responsible for following the cited policy and its existing approval process.')
p('The system boundary includes document ingestion, access enforcement, retrieval, answer generation, citation checking, the API, the chat interface, logging, and evaluation. External inputs are approved files, user credentials, user questions, and evaluation records. An external language-model service receives only the question and authorised evidence needed for generation. The source handbook is acquired offline; the assistant does not browse the live handbook while answering.')
table(['External actor or resource','Input to the assistant','Output from the assistant'],[
('Employee','Credentials; English question; signed token on later requests.','Login result; answer and citations; generic not-found result.'),
('Project operator','Approved documents; department mappings; deployment settings; evaluation set.','Ingestion result; health status; query metrics; evaluation report.'),
('Language-model service','Generated answer and candidate citation identifiers returned to the assistant.','Question, generation instructions, and only authorised evidence sent to the service.'),
('Approved document collection','Markdown files and text-based PDF files.','No source-document edits are produced by the assistant.')],[1.35,2.65,2.8])
p('Table 1 defines the external context. Internal retrieval engines and storage components are described in R4 rather than treated as additional external actors.')
h('2.1.1 System Interfaces',3)
p('The employee interface communicates with the application API for login, questions, and health checks. Offline ingestion reads the configured document folder. Evaluation invokes the same question-answering behaviour under each test role. There is no integration with an organisation payroll system, identity provider, HR application, or live document-management service in version 1.')
h('2.1.2 User Interfaces',3)
p('Employees use a browser-based login form and chat interface. The interface presents the submitted question, the answer, and readable source citations. It also presents not-found, authentication, and service-error outcomes. An operator uses command-line ingestion and evaluation tools. Employees do not receive controls to change their own role or upload documents into the approved collection.')
h('2.1.3 Hardware Interfaces',3)
p('The system has no specialised hardware interface requirements. It uses ordinary host storage, network access, and a browser. It does not control sensors, embedded devices, or external machinery.')
h('2.1.4 Software Interfaces',3)
p('The application exchanges requests and responses with a hosted language-model service. The current design selects Groq, with the model service isolated behind a replaceable client interface. The database, vector engine, keyword scorer, and UI framework are internal design choices rather than customer-mandated software interfaces. Provider-specific versions, authentication details, and request limits belong in the deployment configuration and remain subject to the open interface decisions in Section 6.3.')
h('2.1.5 Communications Interfaces',3)
p('The deployed UI and API exchange structured requests over HTTPS. Authenticated requests carry a bearer token. The backend also needs outbound network connectivity to the configured model provider. Local source-file ingestion requires no network protocol, and no custom device or messaging protocol is required.')
h('2.1.6 Memory Constraints',3)
p('No fixed RAM or storage ceiling has been approved. Resource use depends on the selected embedding model, document collection, indexes, and deployment host. Minimum host specifications and supported concurrency must be established from the first measured baseline; no capacity claim is made before those decisions are recorded.')
h('2.1.7 Operations',3)
p('The first release has separate preparation, serving, and evaluation activities. During preparation, the operator supplies documents, runs ingestion, and prepares the retrieval data. During serving, employees log in and ask questions. Ingestion takes place offline while the question-answering service is stopped. The service is restarted after an update and checks that its stored document information and retrieval indexes agree.')
p('Evaluation is an operator-run activity, not an employee command. Backup frequency, backup location, recovery time, log retention, and deletion arrangements have not yet been approved. These are open operational decisions; the SRS does not promise continuous service during document updates.')
h('2.1.8 Site Adaptation Requirements',3)
p('Before deployment, the operator configures the approved data directory, role assignments, department mapping, provider credentials, storage locations, retrieval settings, and content-logging setting. The demo uses public documents and simulated roles. Replacing it with private company documents requires an agreed access mapping and approval of the external provider data flow. No alteration to the physical work area is required.')
h('2.2 Product Functions',2)
bullets([
'Prepare an approved document collection and preserve its department and source metadata.',
'Authenticate employees and determine their permissions from a backend-verified token.',
'Find relevant passages within the departments permitted for the employee.',
'Produce an evidence-based answer with validated citations, or return a generic not-found outcome.',
'Expose the behaviour through an API and chat interface.',
'Record operational metrics and compare answer quality across retrieval configurations.'
])
h('2.3 User Characteristics',2)
p('Employee users are expected to read English and use a browser, a login form, and a chat box. They do not need knowledge of embeddings, retrieval algorithms, or language-model prompting. The four employee roles represent permission groups; they do not imply different technical skill levels. The operator is the project owner or a technically capable maintainer who can configure the environment and run command-line tools.')
table(['Role','Token value','Permitted departments','Restricted departments'],[
('New Hire','new_hire','onboarding, general','people, finance, engineering'),
('HR','hr','people, onboarding, general','finance, engineering'),
('Finance','finance','finance, onboarding, general','people, engineering'),
('Engineer','engineer','engineering, onboarding, general','people, finance')],[.8,1,2.6,2.4])
p('The operator is an operational responsibility, not an additional employee token role. The public handbook itself has no private access boundary; the project applies this simulated policy for demonstration and testing.')
h('2.4 Constraints',2)
bullets([
'The first release serves one organisation with four employee roles and department-level permissions.',
'Permission checks occur in the backend before evidence reaches answer generation. Client controls and prompts cannot grant access.',
'The serving process uses one API worker. Document ingestion is performed offline.',
'The demo collection is pinned to a recorded handbook commit. Dataset provenance and attribution are retained.',
'External service availability and account limits affect generation and evaluation.',
'Secrets, logs, local datasets, virtual environments, and generated databases are excluded from the public source repository.'
])
h('2.5 Assumptions and Dependencies',2)
p('The requirements are based on the project owner’s stated goals and existing specification; a real enterprise stakeholder interview has not been completed. The source collection is assumed to contain approved and readable policies. The assistant cannot establish that a policy is correct merely because it appears in a document. Source accuracy and classification remain operator responsibilities.')
p('Users are assumed to have a browser and network connectivity to the deployment. The backend requires access to the selected model provider during answer generation. Role assignment is trusted only after backend authentication. A revoked user may retain access until the existing token expires; immediate revocation and enterprise identity synchronisation are deferred.')
p('The original quality thresholds are provisional. Hardware, concurrency, provider configuration, and the scoring method must be recorded before comparing measured results with those thresholds. Changes to these assumptions can require a revision through Section 4.')
h('2.6 Apportioning of Requirements',2)
table(['Release','Included capabilities'],[
('Version 1 baseline','Approved local Markdown and text-based PDF ingestion; four simulated roles; signed-token login; restricted retrieval; grounded answers and citations; API and chat; metrics; evaluation; deployed demo.'),
('Later versions','Enterprise SSO and immediate entitlement checks; finer access groups; background document synchronisation; multiple API workers; optional grounding verification at request time; automatic personal-data redaction; per-client configuration.'),
('Outside the present commitment','Voice, multilingual support, native mobile applications, model fine-tuning, document editing through chat, multi-tenant hosting, and OCR for scanned PDFs.')],[1.4,5.4])
p('An item listed for a later version is not a version 1 acceptance condition. Adding it requires an approved scope change.')

h('3 Specific Requirements',1,True)
p('The requirements in this section apply to the complete planned first release. High-priority requirements are necessary for baseline acceptance. References in brackets identify the project source. Verification identifiers refer to the planned procedures in Section 6.1 and are not claims that those procedures have already passed.')
h('3.1 External Interface Requirements',2)
h('3.1.1 User Interfaces',3)
table(['ID','Required interface behaviour','Verification'],[
('IR-01','The employee interface shall provide credential entry, question entry, answer display, citation display, and visible feedback when authentication or a request fails. [R3 FR-07, FR-11; R4 Section 2]','AT-07, AT-11'),
('IR-02','The employee interface shall send the backend-issued token with question requests and shall not use a selectable role value as authority for access. [R4 Section 4.2]','AT-05, AT-11'),
('IR-03','The interface shall present the same not-found outcome for restricted and missing evidence. [R3 FR-12]','AT-12')],[.65,5.25,.9])
p('No fixed screen resolution, browser support matrix, or formal accessibility conformance level has been approved. These decisions are recorded in Section 6.3. The required interface functions remain applicable while the visual layout is developed.')
h('3.1.2 Hardware Interfaces',3)
p('No direct hardware interfaces are applicable, as stated in Section 2.1.3. Host memory and storage capacity are deployment decisions rather than device-control requirements.')
h('3.1.3 Software Interfaces',3)
table(['Operation','Input and authority','Output and conditions'],[
('POST /auth/login','User credentials; no prior bearer token required.','Authentication result and a signed token on success. Credential field names and error schema remain TBD.'),
('POST /ask','JSON body containing question; a verified bearer token supplies the role.','Answer, citations, and a found flag. The response distinguishes a request failure from an ordinary not-found outcome.'),
('GET /health','No authentication required in the current design.','Service status and consistency status of the retrieval data; public response detail remains TBD.'),
('Ingestion CLI','Configured data directory; department and provenance information.','Prepared document data and retrieval indexes, or an ingestion error. No remote ingestion endpoint is provided.'),
('Evaluation CLI','Question records with id, role, question, expected_answer, source_doc, should_refuse.','Per-question outcomes and aggregate evaluation results.'),
('Model provider','Question, generation instructions, and authorised source passages with chunk identifiers.','Generated answer and proposed citation identifiers, or a provider error.')],[1.2,2.65,2.95])
label('IR-04.','The question API shall derive the role from the verified token rather than from the question body. [R3 FR-11; R4 Section 6] Verification: AT-11.')
label('IR-05.','The model request shall contain no source passage outside the authenticated role’s permitted departments. [R3 FR-05; R4 Section 4.1] Verification: AT-05 with captured provider input.')
p('The question field is text. Maximum question length, whitespace-only input handling, rate limits, credential schema, citation JSON fields, and exact HTTP error codes are open interface decisions. These details must be fixed before the API contract is approved; illustrative payloads are not used as substitutes for a final schema.')
h('3.1.4 Communications Interfaces',3)
label('IR-06.','The deployed user-facing service shall use HTTPS and carry its signed access token as a bearer token on authenticated requests. [R4 architecture and Section 4.2] Verification: AT-11 and deployment inspection.')
p('There is no application-specific network protocol. Provider authentication material belongs in protected deployment configuration. Network timeouts, retry limits, and the exact service-error response are open decisions in Section 6.3.')
h('3.2 System Features and Functional Requirements',2)
p('The feature descriptions below give the actor, trigger, preconditions, processing obligations, and observable results. The original requirement identifiers are preserved so that the new SRS can be traced back to the existing project specification.')

reqs=[]
def requirement(rid,title,statement,source,acceptance,deps='None beyond the stated feature preconditions.'):
    start=len(doc.paragraphs)
    x=p(rid+' '+title,True);x.paragraph_format.keep_with_next=True
    p(statement)
    label('Source and priority.',source+' | High')
    label('Acceptance.',acceptance)
    label('Dependencies.',deps)
    for para in doc.paragraphs[start:-1]:
        para.paragraph_format.keep_with_next=True
        para.paragraph_format.keep_together=True
    reqs.append((rid,title))

h('3.2.1 Document Ingestion and Classification',3)
label('Purpose.','Prepare approved source documents for later retrieval while retaining their origin and permission label.')
label('Actor and trigger.','The operator runs offline ingestion against a configured folder.')
label('Preconditions.','The serving process is stopped, the folder is readable, and source paths can be mapped to known departments.')
label('Sequence and result.','The system reads supported files, extracts text, associates department and source metadata, and prepares passages for storage and retrieval. Invalid or excluded paths are refused. Changed documents replace their earlier derived passages, and deleted documents are removed from the serving collection during the offline update.')
requirement('FR-01','Supported document ingestion',
'The system shall ingest Markdown documents and text-based PDF documents from a configured local data folder.',
'R3 FR-01; NFR-05; R4 Section 3.1',
'AT-01: ingest one readable Markdown fixture and one text-based PDF fixture; confirm that their body text is available in the prepared collection and remains associated with the source file. Scanned PDFs are outside the required set.')
requirement('FR-02','Department and source metadata',
'The system shall associate each ingested document and its derived chunks with a department and source metadata that supports access enforcement and source citation.',
'R3 FR-02; R4 Section 5; R5 department mapping',
'AT-02: verify department, document identity, source path, checksum, title, and section or page metadata for prepared fixtures; reject a source that cannot be assigned an approved department.','FR-01; BR-01 to BR-03; DR-01 to DR-06.')
label('Exception handling.','An unmapped, excluded, or malformed source path must not receive a default department. Processing errors must not make unauthorised or unclassified content available to employees. Rules for corrupt PDFs and partial-ingestion reporting remain open under OD-04.')

h('3.2.2 Authentication and Role Enforcement',3)
label('Purpose.','Establish an employee’s identity and apply the access policy before source content can be supplied for generation.')
label('Actor and trigger.','An employee submits credentials and later asks a question using the issued token.')
label('Preconditions.','A valid demo account has an assigned role, and the backend can verify its token.')
label('Sequence and result.','Successful authentication produces a signed token. Each question request is checked by the backend. The verified role selects the permitted departments in Section 2.3. An unknown role does not receive a fallback permission set. The current design expires tokens after 30 minutes; immediate revocation is outside version 1.')
requirement('FR-11','Signed-token authentication',
'The system shall authenticate an employee and determine the employee’s role from a signed token verified by the backend.',
'R3 FR-11; R4 Section 4.2',
'AT-11: valid credentials produce a usable signed token; altered, invalid, and expired tokens cannot authorise a question. A role supplied in the question body cannot change the role determined by the token.','IR-02, IR-04, IR-06; the role policy in Section 2.3.')
requirement('FR-05','Retrieval within permitted departments',
'The system shall supply retrieval evidence only from documents within the authenticated employee’s permitted departments.',
'R3 FR-05; R4 Sections 4.1 and 4.7',
'AT-05: for each role, inspect retrieved chunks and the complete model request; every passage must belong to an allowed department. A search dominated by disallowed documents must still consider permitted candidates before selecting the final results.','FR-02, FR-11; BR-01 and BR-04; IR-05.')
requirement('FR-12','Indistinguishable restricted and missing outcomes',
'The system shall return the same generic not-found reply when requested evidence is restricted and when it does not exist in the approved collection.',
'R3 FR-12; R4 Section 4.4',
'AT-12: ask equivalent questions about a known restricted source and a non-existent source as the same role; compare the public found flag, answer message, and citations. The not-found result must not disclose the restricted title or its existence.','FR-05, FR-06. Exact reply wording is fixed under OD-04; no timing-equivalence guarantee is asserted.')

h('3.2.3 Question Answering and Evidence Sufficiency',3)
label('Purpose.','Answer questions from available evidence while refusing unsupported conclusions.')
label('Actor and trigger.','An authenticated employee submits an English question.')
label('Preconditions.','The service is ready, the token is valid, and a consistent approved collection is available.')
label('Sequence and result.','The system retrieves permitted passages and evaluates whether they are relevant enough to support generation. If the evidence is insufficient, it returns not found without inventing an answer. Otherwise, the model receives the question and permitted evidence; its answer then passes citation validation before release.')
requirement('FR-03','Answers grounded in retrieved documents',
'The system shall answer a natural-language question using the retrieved document content as its source of factual evidence.',
'R3 FR-03; R4 Section 3.2',
'AT-03: submit factual and multi-section questions with expected answers derived from approved sources; assess answer correctness and source support using the evaluation method defined for NFR-02.','FR-01, FR-02, FR-05, FR-11.')
requirement('FR-06','Refusal when evidence is insufficient',
'The system shall return a not-found outcome when the permitted retrieved evidence does not support an answer.',
'R3 FR-06; R4 Section 4.3',
'AT-06: below-threshold evidence stops generation; an insufficient-context result from generation also becomes not found. Questions outside the collection must not produce a factual answer based solely on model knowledge.','FR-05; BR-05. Threshold calibration and the exact message remain controlled open decisions.')
requirement('FR-10','Resistance to instruction overrides',
'The system shall preserve its access rules and evidence-based answering behaviour when a question or retrieved document contains instructions that attempt to override them.',
'R3 FR-10; R4 Section 4.6',
'AT-10: test instructions embedded in questions and document passages that request another role, restricted sources, unsupported answers, or invented citations. Inspect both the public result and the model input; no restricted passage may be supplied.','FR-05, FR-06, FR-11, FR-13. Passing a defined test set is evidence for that set, not a universal security proof.')

h('3.2.4 Citation Validation and Answer Delivery',3)
label('Purpose.','Make a released answer traceable to the evidence actually supplied for its generation.')
label('Actor and trigger.','The generator produces an answer and proposed source identifiers.')
label('Preconditions.','The candidate evidence has passed the access filter and is available for citation lookup.')
label('Sequence and result.','The system checks proposed identifiers against the passages supplied to the model. Valid sources become readable citations. Invalid identifiers are removed. If no valid citation remains, the system suppresses the unsupported answer and returns not found; it does not make another generation call solely to repair citations.')
requirement('FR-04','Readable source citations',
'The system shall include at least one citation identifying the document title and section or PDF page in every substantive answer returned to an employee.',
'R3 FR-04; R4 Section 5',
'AT-04: each successful found=true response includes a citation with a document title and section or page that resolves to supplied evidence. A not-found response is not a substantive answer and requires no invented citation.','FR-02, FR-03, FR-13.')
requirement('FR-13','Citation identifier validation',
'The system shall validate every proposed citation against the chunks supplied to the model, discard invalid identifiers, and return not found when no valid citation remains.',
'R3 FR-13; R4 Section 4.8',
'AT-13: all-valid identifiers are retained; mixed valid and invalid identifiers retain only valid entries; all-invalid or empty identifiers produce not found. Confirm that citation repair does not trigger another model call.','FR-02, FR-05, FR-06; BR-06.')
requirement('FR-07','API and chat access',
'The system shall make question answering available through a REST API and a browser-based chat interface.',
'R3 FR-07; R4 Section 6',
'AT-07: authenticate and ask the same supported question through each interface; confirm equivalent access restrictions, answer outcome, and citation requirements. Confirm that ingestion and evaluation remain operator tools rather than employee API functions.','FR-03 to FR-06, FR-11 to FR-13; IR-01 to IR-06.')

h('3.2.5 Query Logging and Evaluation',3)
label('Purpose.','Measure behaviour, investigate retrieval outcomes, and compare configurations without storing question content by default in client deployments.')
label('Actor and trigger.','A query completes, or the operator starts an evaluation run.')
label('Preconditions.','Logging settings are configured; evaluation records identify a role and expected outcome.')
label('Sequence and result.','A completed query records operational metrics. The evaluation runner submits the acceptance questions using their assigned roles, scores the results, and reports quality, refusal, safety, timing, and cost measures. Configuration comparisons use the same recorded question set and dataset revision.')
requirement('FR-08','Query metrics with controlled content logging',
'The system shall record query role, retrieved chunk identifiers, relevance scores, latency, and token count; it shall store question and answer text only when content logging is enabled.',
'R3 FR-08; R4 Sections 4.10 and 7',
'AT-08: with content logging disabled, required metric fields are present and question/answer text is absent. With it enabled, text may be recorded. Logged role values are bounded and sanitised. The client-deployment default is content logging off.','FR-07; DR-07. Retention and limits are fixed under OD-02 and OD-04.')
requirement('FR-09','Repeatable evaluation report',
'The system shall provide an evaluation script that runs the defined test-question set and reports the resulting scores.',
'R3 FR-09; R4 Section 7',
'AT-09: provide a valid evaluation file, run it as each recorded role, and obtain traceable per-question and aggregate results. Record the dataset revision, configuration, scoring method, and unresolved or failed questions.','FR-03 to FR-08, FR-10 to FR-13; the acceptance set in Section 6.1.')
p('The design calls for bounded evaluation concurrency, retries with backoff on provider rate limits, and cached results that permit resuming an interrupted run. Their numeric settings and cache invalidation details remain configuration decisions. They must not permit a cached answer from one role or configuration to bypass the required access policy.')

h('3.2.6 Business Rules',3)
table(['ID','Rule','Source'],[
('BR-01','The only employee roles are new_hire, hr, finance, and engineer. Allowed departments are exactly those in Section 2.3. Unknown roles fail closed.','R3 Section 3; R6 config.py'),
('BR-02','Department assignment follows the longest matching source-path prefix. An explicit file mapping overrides its parent folder. An unmapped source is refused.','R5; R6 departments.py'),
('BR-03','Paths are relative and use forward slashes. Segments allow ASCII letters, digits, period, underscore, and hyphen. Traversal, empty segments, trailing dots or spaces, and excluded prefixes are refused.','R5 path rules'),
('BR-04','Access restriction applies to vector and keyword results before final candidate selection and before any passage reaches the model.','R4 Sections 4.1, 4.7'),
('BR-05','The not-found threshold uses reranker relevance when enabled, otherwise best vector similarity. A fused RRF score is not the sufficiency measure. Each configuration has a calibrated threshold.','R4 Section 4.9'),
('BR-06','Citations identify only chunks supplied for the current generation. Invalid identifiers are removed; no valid citation means not found.','R3 FR-13; R4 Section 4.8'),
('BR-07','A changed or deleted document loses its obsolete chunks during offline ingestion. Unchanged documents need not be embedded again.','R4 Section 3.1'),
('BR-08','The demo selection contains 190 pages: engineering 60, finance 35, people 50, onboarding 5, and general 40. Each selected page has at least 150 real words by the documented selection method.','R5 selection rules')],[.65,4.8,1.35])
p('The excluded source prefixes are finance/readmes and finance/jihu-support. Prefix exclusions are checked without regard to case; approved mapping matches are case-sensitive. The complete mapping is maintained in R5 and the implementation mapping table, rather than copied into multiple access-policy sources.')

h('3.3 Performance Requirements',2)
p('These targets retain the existing project values. NFR-01, NFR-02, and NFR-04 are provisional until the test environment and scoring choices in Section 6.3 are approved. A measured result must identify those conditions; an unqualified claim of “fast” or “accurate” is not acceptable evidence.')
table(['ID and priority','Requirement and target','Verification and qualification'],[
('NFR-01\nMedium','The system shall complete at least 95% of answer requests within 8 seconds under the agreed acceptance workload.','PT-01: measure elapsed time from API receipt of a valid question to the complete answer response. Record hardware, dataset, configuration, concurrent users, request count, and warm-up policy. Concurrency and the final protocol remain OD-01. Source: R3 NFR-01.'),
('NFR-02\nMedium','The evaluated system shall achieve faithfulness of at least 0.85 and answer correctness of at least 75% on the agreed test questions.','QT-01: retain per-question scores and aggregate results. The scoring rubric, scoring model if used, treatment of refusals, and denominator must be approved under OD-05 before a pass can be declared. Source: R3 NFR-02.'),
('NFR-04\nMedium','The average language-model charge shall not exceed USD 0.02 per evaluated question.','CT-01: divide total billed generation cost by the evaluated question count and retain token usage and the price basis used for that run. State whether evaluation-judge calls are excluded. The accounting boundary remains OD-05. Source: R3 NFR-04.')],[1.05,2.7,3.05])
p('The first demonstration uses 190 selected documents. This number is a reproducible dataset baseline, not a claim about maximum production capacity. Simultaneous-user capacity, throughput, memory ceilings, and maximum source-file size remain undecided under OD-01 and OD-04. A single API worker is an operating constraint, not a one-user limit.')

h('3.4 Logical Database Requirements',2)
p('The logical information model is independent of the selected storage engine. Each document has one source path, one department, and a checksum for the ingested revision. A document produces multiple chunks. Citations and query records refer to chunks, while retrieval indexes are derived from the authoritative chunk collection.')
table(['Entity','Required information','Relationship or constraint'],[
('Document','document_id, source_path, title, department, checksum.','document_id identifies a source document; a document has one approved department and zero or more derived chunks.'),
('Chunk','chunk_id, document_id, department, text, citation title, section or page locator.','chunk_id is unique. document_id links to its source. Department agrees with the document; locator supports citation.'),
('Role policy','Role name and permitted department set.','A single access table governs the four employee roles; chunks do not carry a second conflicting role list.'),
('Manifest record','source_path, department, sha256.','Source paths are unique and sorted. The checksum is a 64-character lowercase hexadecimal SHA-256 value.'),
('Query record','Role, retrieved chunk IDs, relevance scores, final chunk IDs, latency, token count; optional question and answer text.','Identifiers refer to the evidence used by that query. Content fields depend on the content-logging switch.'),
('Evaluation question','id, role, question, expected_answer, source_doc, should_refuse.','The expected answer and refusal flag are written against the pinned source collection.'),
('Evaluation result','Question reference, observed outcome, configuration and dataset revision, reported scores and measurements.','Results remain traceable to the question set, source revision, and scoring conditions.')],[1.05,2.7,3.05])
table(['ID','Integrity or lifecycle rule','Verification'],[
('DR-01','The system shall retain document identity and source metadata needed to trace every served chunk to its source.','AT-02'),
('DR-02','The system shall reject chunks labelled with a department outside the approved department set.','AT-02'),
('DR-03','The system shall use a single authoritative chunk collection from which retrieval indexes are built or updated.','DT-01'),
('DR-04','The system shall remove obsolete chunks for a changed or deleted document from the serving collection during offline ingestion.','DT-01'),
('DR-05','The system shall check chunk counts and checksums across the stored collection and derived indexes at startup and refuse to serve when they disagree.','DT-01'),
('DR-06','The demo manifest shall use source checksums calculated after CRLF line endings are normalised to LF.','DT-02'),
('DR-07','The system shall omit question and answer text from query records when content logging is disabled.','AT-08')],[.7,5.05,1.05])
p('Rules DR-01 and DR-02 elaborate R3 FR-02 and R4 Section 5; DR-03 to DR-05 elaborate R4 Section 3.1; DR-06 comes from R5; DR-07 elaborates R3 FR-08. In the current source-based identity convention, document_id is the source path and chunk_id adds a numbered suffix such as #000. The persistence representation remains in the design document.')
p('Documents remain in the approved collection until the operator changes or removes them through offline ingestion. Log retention, deletion scheduling, evaluation-record retention, and backups are not yet baselined. OD-02 records these decisions; version 1 does not imply indefinite retention. No real-user password storage schema has been approved, so this SRS does not prescribe a credentials table.')

h('3.5 Design Constraints',2)
p('The SRS records established constraints without turning every selected library into a requirement. The planned stack in R4 includes Python 3.12, FastAPI, Streamlit, SQLite, ChromaDB, BM25, a local embedding model, a cross-encoder reranker, and Groq. These selections may be revised through design review when the required observable behaviour is preserved.')
table(['ID','Constraint','Basis and verification'],[
('DC-01','Version 1 shall run with a single API worker and perform ingestion offline while serving is stopped.','R4 Sections 3.1 and 8; inspect deployment settings and demonstrate the update procedure.'),
('DC-02','The backend shall enforce access before evidence reaches generation; client-supplied role choices shall not grant authority.','R4 Sections 4.1 and 4.2; AT-05 and AT-11.'),
('DC-03','Ingestion and evaluation shall remain operator-run command-line activities without remote document-modification endpoints.','R4 Section 6; inspect routes and operator commands.'),
('DC-04','The demo shall identify the pinned GitLab Handbook source revision and retain its attribution record.','R5 and R6 README; DT-02 and documentation inspection.')],[.7,4.2,1.9])
h('3.5.1 Standards Compliance',3)
p('This document uses the supplied course guidance for requirement quality and adapts the supplied template’s feature-based outline. The data-source record identifies the handbook licence as CC BY-SA 4.0. Dataset attribution is preserved in the project documentation. No additional regulatory certification, accessibility certification, or enterprise security certification has been established for this student demonstration.')
p('The product is a document-information assistant. It is not specified as a medical, legal, payroll, or financial transaction authority. Any new domain-specific compliance obligations must be identified and approved before real deployment in that domain.')

h('3.6 Software System Attributes',2)
h('3.6.1 Reliability',3)
p('Reliability is evaluated through evidence sufficiency, validated citations, repeatable results, and consistency of the serving data. FR-06 prevents an answer when evidence is insufficient; FR-13 suppresses an answer with no valid citation; DR-05 prevents serving an inconsistent index. QT-01 assesses whether generated claims remain supported by the cited material.')
p('No numerical mean-time-between-failure target has been agreed. A provider or infrastructure failure is a service error rather than proof that the documents contain no answer. The exact timeout, retry, recovery, and public error behaviour must be settled under OD-04 before operational acceptance.')
h('3.6.2 Availability',3)
p('The demonstration is available when the host and required provider are running. Offline document updates require a service stop and restart. A health endpoint exposes service status in the current design. No 24-hour service commitment, uptime percentage, recovery-time objective, or recovery-point objective has been approved. OD-02 covers backup and recovery expectations; OD-06 covers the deployment host.')
h('3.6.3 Security and Privacy',3)
table(['ID and priority','Requirement','Verification'],[
('NFR-03\nHigh','The system shall disclose no restricted document content across the defined access-control acceptance questions. [R3 NFR-03]','ST-01: inspect answers, citations, retrieved evidence, and model payloads for all role-denial cases. Any leak fails acceptance for the test set.'),
('NFR-07\nHigh','No API key, signing secret, password, or other deployment secret shall be committed to the project repository. [R3 NFR-07]','ST-02: inspect tracked files and configuration examples for secrets; confirm sensitive local artifacts are excluded. Record the inspection scope.')],[1.05,3.15,2.6])
p('Security also depends on FR-05, FR-10, FR-11, FR-12, and FR-13. Public not-found results must not expose the existence of restricted files. Question content does not have authority to alter the role policy. Development logs may contain public demo text when content logging is enabled; client deployments default to content logging off. Automatic personal-data masking is deferred.')
p('A valid token expires after 30 minutes in the current design. Immediate access revocation is not guaranteed during that interval. Signing algorithm, credential storage, account provisioning, sanitisation bounds, provider handling of submitted content, and any browser session-storage decisions remain OD-03 and OD-04.')
h('3.6.4 Maintainability',3)
p('The established design isolates model and vector-store implementations behind interfaces and keeps the role-to-department mapping in one place. Changing a provider must preserve the SRS behaviour and pass the affected acceptance procedures. A replacement must not relax access rules, alter not-found semantics, or invalidate citation checks. Source, tests, requirements, and design are maintained under version control, with traceability recorded in Section 6.2.')
p('No numerical code-complexity, code-coverage, or repair-time threshold is established. Maintainability is reviewed against the documented interfaces and whether the same acceptance checks remain usable after a component change.')
h('3.6.5 Portability and Reproducibility',3)
table(['ID and priority','Requirement','Verification'],[
('NFR-06\nMedium','The system shall be set up from its README using Docker in less than 30 minutes under the agreed setup conditions. [R3 NFR-06]','RT-01: use a clean supported host and follow the README. Record host, network, downloads, start/end times, and deviations. Conditions remain OD-06.'),
('NFR-05\nHigh','The system shall support Markdown and text-based PDF source documents. Scanned PDFs are not required. [R3 NFR-05]','AT-01: verify both required formats. Record parser limitations and reject or report unsupported content according to the approved ingestion contract.')],[1.05,3.15,2.6])
p('Development uses Windows and PowerShell, while project tests are intended to remain compatible with Linux. A complete operating-system and browser support matrix has not been agreed. Portability does not imply support for every operating system or arbitrary resource-limited host.')
h('3.6.6 Usability Accessibility and Scalability',3)
p('IR-01 requires the basic login, question, answer, citation, and error interactions. The product exposes ordinary user tasks rather than internal retrieval settings. No measured usability threshold or accessibility conformance level has been approved. OD-07 requires an agreed browser and accessibility scope before such claims are included in acceptance.')
p('The one-worker deployment and department-level access policy define the first release. Multi-worker hosting, fine-grained entitlements, and multiple organisations are future capabilities. OD-01 requires capacity measurement before stating supported concurrency or throughput.')
h('3.7 Organisation of Specific Requirements',2)
p('Section 3.2 uses the feature organisation offered by the supplied template: ingestion, authentication and permissions, answering, citation delivery, and evaluation. Each feature identifies its purpose and stimulus-response sequence before listing its requirements. The template’s alternatives by mode, user class, object, stimulus, response, functional hierarchy, and use case are organising options rather than additional required product functions. They are therefore not duplicated as empty requirement sections.')
h('3.8 Additional Comments',2)
p('A not-found result is an ordinary question-answering outcome. It is distinct from invalid authentication, invalid input, or an unavailable provider. A valid citation proves that a referenced passage was supplied to generation; it does not by itself prove every sentence in the answer is supported. Faithfulness evaluation is required for that assessment.')
p('Existing code and tests demonstrate implementation progress, but this specification covers the completed intended release. No implementation or evaluation result is presented here as a completed acceptance record. Outstanding decisions in Section 6.3 must be resolved before claiming the affected requirements have a fully approved acceptance baseline.')

h('4 Change Management Process',1,True)
h('4.1 Submitting a Change',2)
p('A proposed change is recorded in writing as a repository issue or equivalent change record. The record identifies the affected requirement IDs, proposed wording, reason, source of the request, and expected effect on scope. A conversation or informal suggestion does not itself change the approved SRS.')
h('4.2 Impact Assessment',2)
p('The project owner checks the effect on roles, source data, interfaces, dependencies, security, quality targets, implementation, and verification. Changes to access rules or provider data flow require explicit review of the affected acceptance cases. New capacity or timing targets require measured evidence or an agreed test protocol rather than an arbitrary number.')
h('4.3 Decision and Approval',2)
p('For this individual project, Abdul Samad maintains the specification and records the accept, reject, or defer decision. Academic approval is recorded by the course reviewer when required. A proposed change affecting a reviewed baseline is submitted to the reviewer before that baseline is treated as approved. These procedures take effect when this document is approved.')
h('4.4 Updating the Baseline',2)
p('An accepted change updates the relevant requirement, its traceability row, acceptance procedure, design references, open-decision record, and revision history. Existing identifiers are retained where the obligation is unchanged; withdrawn identifiers remain recorded and are not reused for unrelated behaviour. The updated document receives a new version. Approved versions are retained so that a reviewer can identify which requirements governed an implementation or test run.')
table(['Change record field','Required entry'],[
('Change identifier','Unique reference for the request.'),('Affected requirements','FR, NFR, supporting rule, interface, and acceptance IDs.'),('Requested change and reason','Current wording, proposed wording, and rationale.'),('Impact','Scope, data, security, design, tests, cost, and schedule implications.'),('Decision','Accepted, rejected, or deferred, with the decision maker and date.'),('Baseline reference','Document version and implementation or test references after approval.')],[1.7,5.1])
p('No change approvals are implied by the presence of a proposed rule or provisional target in this first version.')

h('5 Document Approvals',1,True)
p('Approval confirms that the identified version is an agreed basis for the stated purpose. Signatures and approval dates are left unfilled until the named person reviews the document. An academic review does not represent approval by a real enterprise customer.')
table(['Role','Name','Decision','Signature','Date'],[
('Author and project owner','Abdul Samad','Pending','________________','____________'),
('Course reviewer','Sir Fahim Muhammad Khan','Pending','________________','____________')],[1.3,1.7,.85,1.7,1.25])
h('5.1 Approval Conditions',2)
p('The reviewer checks scope, requirement clarity, role policy, source traceability, and verification coverage. Open decisions may remain in an academically reviewed draft, but an operational acceptance baseline must identify which have been resolved and which remain explicitly deferred. No unapproved provisional condition is reported as a passed acceptance criterion.')
h('5.2 Release Acceptance Record',2)
p('Document approval and product acceptance are separate. A completed product acceptance record identifies the tested product revision, dataset revision, configuration, environment, dates, executed procedures, results, unresolved defects, and final acceptance decision. That record will be completed after the planned system is implemented and evaluated.')

h('6 Supporting Information',1,True)
p('Sections 6.1 and 6.2 define verification and traceability for this SRS. Section 6.3 is the controlled register of unresolved decisions. Sections 6.4 to 6.6 supply provenance, current implementation context, and an index. The implementation context is informative; it does not replace the requirements in Section 3.')
h('6.1 Verification and Acceptance Plan',2)
p('Verification uses tests, inspection, analysis, and demonstration. A test result records the observed outcome rather than merely naming the test. The planned core question set contains 30 records in eval/questions.csv. Questions and expected answers are written from the pinned handbook pages, not invented independently of the collection.')
table(['Question category','Count','Purpose'],[
('Simple factual',12,'Check direct retrieval, correctness, and citation support.'),('Multi-section or table',6,'Check answers requiring evidence from several passages.'),('Restricted access',6,'Check refusal and absence of restricted evidence.'),('Outside the documents',6,'Check not-found behaviour when no answer exists.')],[1.75,.55,4.5])
p('The 30 questions are the original target set. They are supplemented by authentication, prompt-injection, citation-validation, data-integrity, and interface fixtures. Access-policy fixtures exercise all four roles and both allowed and denied department pairs; the six core refusal questions alone do not establish exhaustive policy coverage.')
table(['Procedure','Required evidence','Coverage'],[
('AT-01','Markdown and text-based PDF ingestion outputs; source traceability; unsupported-format record.','FR-01, NFR-05'),
('AT-02','Per-document and per-chunk metadata; valid and invalid path cases; mapping overrides and exclusions.','FR-02; BR-02, BR-03; DR-01, DR-02'),
('AT-03','Factual and multi-section answers with expected-answer comparisons.','FR-03'),
('AT-04','Successful answer citations resolved to document title and section or page.','FR-04'),
('AT-05','Each role’s allowed and denied retrieval cases; captured evidence and model inputs; keyword-filter ordering case.','FR-05; BR-01, BR-04; IR-02, IR-05; DC-02'),
('AT-06','Below-threshold, outside-collection, and model-insufficient-context outcomes.','FR-06; BR-05'),
('AT-07','Authenticated API and browser chat demonstration with matching access behaviour.','FR-07; IR-01'),
('AT-08','Logs with content logging both on and off; required metric fields; sanitisation.','FR-08; DR-07'),
('AT-09','Evaluation file, configuration, source revision, per-question results, and aggregate report.','FR-09'),
('AT-10','Question and document prompt-injection fixtures; public results and provider input inspection.','FR-10'),
('AT-11','Valid, invalid, expired, altered, and role-tampered authentication cases; HTTPS inspection.','FR-11; IR-02, IR-04, IR-06'),
('AT-12','Restricted and missing source outcomes compared for public response equivalence.','FR-12; IR-03'),
('AT-13','All-valid, mixed, all-invalid, and empty citation sets; generation-call count.','FR-13; BR-06')],[.8,4.1,1.9])
table(['Procedure','Required evidence','Coverage'],[
('PT-01','Complete latency samples and percentile calculation under the approved environment and workload.','NFR-01'),
('QT-01','Faithfulness and correctness rubric, per-question scores, aggregate calculations, and refusal treatment.','NFR-02'),
('ST-01','Zero restricted-content disclosures across the recorded access-control fixtures.','NFR-03'),
('CT-01','Tokens, run cost, pricing basis, question count, and mean cost calculation.','NFR-04'),
('RT-01','Clean-host setup transcript, conditions, elapsed time, and README deviations.','NFR-06'),
('ST-02','Repository-secret inspection and excluded-artifact review with a stated inspection scope.','NFR-07'),
('DT-01','Initial load, unchanged update, changed document, deletion, and deliberately inconsistent-index cases.','BR-07; DR-03 to DR-05; DC-01'),
('DT-02','Manifest uniqueness, order, quotas, mappings, checksum format, local-file agreement, and attribution record.','BR-08; DR-06; DC-04'),
('DI-01','API route inventory and operator-tool demonstration showing no remote ingestion or evaluation route.','DC-03')],[.8,4.1,1.9])
label('Acceptance rule.','A required behaviour passes only when its stated procedure produces the expected evidence. Any restricted-content leak fails the safety criterion. Timing, quality, cost, and setup targets cannot receive a final pass until their open measurement conditions are approved. Failed or unexecuted procedures remain visible in the acceptance report.')
label('Comparison control.','Vector-only, hybrid, and hybrid-plus-reranker runs use the same pinned dataset and question set. Each run records its own threshold and configuration. False refusals and false answers are reported so a permissive threshold is not mistaken for better overall quality.')

h('6.2 Requirements Traceability Matrix',2)
p('The matrix links each original requirement to its project source, its location in this SRS, and planned verification. Supporting rules are also traced so that security, data, and interface obligations do not disappear between specification and testing.')
table(['ID','Source','SRS section','Verification'],[
('FR-01','R3 FR-01; R4 3.1','3.2.1','AT-01'),('FR-02','R3 FR-02; R4 5; R5','3.2.1','AT-02'),
('FR-03','R3 FR-03; R4 3.2','3.2.3','AT-03, QT-01'),('FR-04','R3 FR-04; R4 5','3.2.4','AT-04'),
('FR-05','R3 FR-05; R4 4.1, 4.7','3.2.2','AT-05, ST-01'),('FR-06','R3 FR-06; R4 4.3','3.2.3','AT-06'),
('FR-07','R3 FR-07; R4 6','3.2.4','AT-07'),('FR-08','R3 FR-08; R4 4.10, 7','3.2.5','AT-08'),
('FR-09','R3 FR-09; R4 7','3.2.5','AT-09'),('FR-10','R3 FR-10; R4 4.6','3.2.3','AT-10'),
('FR-11','R3 FR-11; R4 4.2','3.2.2','AT-11'),('FR-12','R3 FR-12; R4 4.4','3.2.2','AT-12'),
('FR-13','R3 FR-13; R4 4.8','3.2.4','AT-13'),('NFR-01','R3 NFR-01','3.3','PT-01'),
('NFR-02','R3 NFR-02','3.3','QT-01'),('NFR-03','R3 NFR-03','3.6.3','ST-01'),
('NFR-04','R3 NFR-04','3.3','CT-01'),('NFR-05','R3 NFR-05','3.6.5','AT-01'),
('NFR-06','R3 NFR-06','3.6.5','RT-01'),('NFR-07','R3 NFR-07','3.6.3','ST-02')],[1,2.9,1.1,1.8])
table(['Supporting IDs','Source and purpose','SRS section','Verification'],[
('IR-01 to IR-03','R3 FR-07, FR-11, FR-12; employee interface and outcome presentation.','3.1.1','AT-07, AT-11, AT-12'),
('IR-04 to IR-06','R4 2, 4.1, 4.2, 6; token authority, model input, and transport.','3.1.3 to 3.1.4','AT-05, AT-11'),
('BR-01','R3 roles; R6 config.py; one permitted-department policy.','3.2.6','AT-05, AT-11'),
('BR-02 to BR-03','R5; R6 departments.py; classification and path validation.','3.2.6','AT-02'),
('BR-04','R4 4.1, 4.7; access before final selection and generation.','3.2.6','AT-05'),
('BR-05','R4 4.9; meaningful evidence threshold per configuration.','3.2.6','AT-06'),
('BR-06','R3 FR-13; R4 4.8; citation membership.','3.2.6','AT-13'),
('BR-07','R4 3.1; changed and deleted source lifecycle.','3.2.6','DT-01'),
('BR-08','R5; pinned selection and department quotas.','3.2.6','DT-02'),
('DR-01 to DR-02','R3 FR-02; R4 5; valid identity and metadata.','3.4','AT-02'),
('DR-03 to DR-05','R4 3.1; authoritative chunks and index integrity.','3.4','DT-01'),
('DR-06','R5; portable source fingerprint.','3.4','DT-02'),
('DR-07','R3 FR-08; content logging control.','3.4','AT-08'),
('DC-01','R4 3.1, 8; one worker and offline ingestion.','3.5','DT-01; deployment inspection'),
('DC-02','R4 4.1, 4.2; backend access enforcement.','3.5','AT-05, AT-11'),
('DC-03','R4 6; operator-only document preparation.','3.5','DI-01'),
('DC-04','R5; R6 README; provenance and attribution.','3.5','DT-02')],[1.2,3.15,1.05,1.4])

h('6.3 Open Decisions',2)
p('The following values and policies are intentionally unassigned. They are not assumed to be zero, unlimited, or already approved. The project owner records each decision and updates affected requirements before the corresponding acceptance result is finalised.')
table(['ID','Decision required','Affected items','Closure evidence'],[
('OD-01','Acceptance hardware; RAM and storage needs; simultaneous users; throughput; latency request count, warm-up, and run conditions.','NFR-01; Sections 2.1.6 and 3.3','Approved benchmark protocol and first baseline results.'),
('OD-02','Log and evaluation-record retention; deletion process; backup frequency and location; recovery objectives and procedure.','FR-08; Sections 2.1.7, 3.4, 3.6.2','Documented operating and retention policy.'),
('OD-03','Credential schema and storage; account and role provisioning; signing method; provider/model version; private-data handling approval.','FR-11; Sections 3.1.3 and 3.6.3','Reviewed authentication and deployment configuration specification.'),
('OD-04','Question and file-size limits; invalid input; corrupt/unsupported files; exact JSON and error schemas; generic reply wording; timeouts; retry bounds; log sanitisation limits.','FR-01, FR-06 to FR-08, FR-12; IR-01 to IR-06','Approved interface and error contract, with boundary-case tests.'),
('OD-05','Quality rubric; scoring tool/version; correctness denominator; refusal treatment; calibrated thresholds; cost-accounting scope and price basis.','NFR-02, NFR-04; FR-06, FR-09','Frozen evaluation protocol and scored baseline report.'),
('OD-06','Demo hosting target; supported operating systems; Docker prerequisites; whether setup timing includes downloads and credentials.','FR-07; NFR-06','Deployment record and clean-host setup protocol.'),
('OD-07','Supported browsers, screen assumptions, keyboard use, accessibility scope, and any usability task threshold.','IR-01; Section 3.6.6','Agreed interface acceptance checklist.')],[.65,3.2,1.25,1.7])
p('No real enterprise customer has approved this specification. Any later customer interview may add requirements or change priorities through the process in Section 4. The academic reviewer’s signature is also pending in Section 5.')

h('6.4 Data Provenance and Attribution',2)
p('The demonstration uses the public GitLab Handbook at commit 2c77a1f5b8c8a80cb7b5151ff11cfad84f98bfbb. The project data-source record gives a download date of 11 October 2026. Selected pages are copied without modification into the department-organised local data folder; that local collection is excluded from Git.')
table(['Department','Selected pages'],[('engineering',60),('finance',35),('people',50),('onboarding',5),('general',40),('Total',190)],[4.8,2])
p('Selection uses a minimum of 150 real words after front matter, link-only lines, and template lines are excluded. Deterministic round-robin selection across sibling folders spreads the sample across topics; the final manifest is sorted by source path. Checksums normalise CRLF to LF. These records support reproducible evaluation rather than making claims about the freshness of the live handbook.')
p('Dataset attribution: GitLab Handbook, GitLab, https://handbook.gitlab.com, source repository https://gitlab.com/gitlab-com/content-sites/handbook, licensed as recorded by the project under Creative Commons Attribution-ShareAlike 4.0 International, https://creativecommons.org/licenses/by-sa/4.0/. The SRS describes the application and does not reproduce handbook policy text.')

h('6.5 Implementation Context',2)
p('The reviewed application source contains the immutable role policy, department-path classification, deterministic dataset selection and manifest creation, and Markdown chunking with front-matter removal. The title-extraction work is in progress. The current chunk record does not yet contain all metadata required by the completed system.')
p('The authenticated API, chat UI, chunk database, retrieval indexes, reranking, generation, citation delivery, query logging, evaluation runner, Docker setup, and hosted demonstration remain planned portions of the complete first release. This SRS states their required behaviour without marking them as delivered. Project code and tests were used as context; no source changes are part of this document deliverable.')

h('6.6 Subject Index',2)
table(['Subject','Location'],[
('Access permissions and roles','2.3; 3.2.2; BR-01 and BR-04'),('Authentication and tokens','3.1.3; 3.2.2; 3.6.3'),
('Citations and source support','3.2.3 to 3.2.4; AT-04 and AT-13'),('Data ingestion and classification','3.2.1; 3.2.6; 3.4'),
('Dataset provenance','1.4; 6.4'),('Document approval and revisions','Document Control; Sections 4 and 5'),
('Evaluation and acceptance','3.2.5; 6.1 to 6.2'),('Logging and privacy','FR-08; DR-07; 3.6.3'),
('Not-found behaviour','FR-06; FR-12; BR-05; 3.8'),('Open decisions','6.3'),
('Performance and cost','3.3'),('Portability and setup','3.6.5'),('Release boundary and future scope','1.2; 2.6')],[2.9,3.9])

doc.core_properties.title='Enterprise RAG Assistant Software Requirements Specification'
doc.core_properties.subject='Complete planned version 1 requirements'
doc.core_properties.author='Abdul Samad'
doc.core_properties.keywords='SRS, Enterprise RAG Assistant, requirements, 243203624'
doc.core_properties.comments=''
doc.core_properties.version='1.0'
settings=doc.settings.element
for st in styles:
    for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    for f in list(st.element.iter(qn('w:rFonts'))):
        for attr in list(f.attrib):
            if 'theme' in attr.lower():del f.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+attr),'Times New Roman')
    for e in list(st.element.iter(qn('w:color'))):
        for attr in list(e.attrib):
            if attr!=qn('w:val'):del e.attrib[attr]
        e.set(qn('w:val'),'000000')
for node in list(settings):
    if node.tag in [qn('w:updateStyles'),qn('w:attachedTemplate')]:settings.remove(node)
for para in doc.paragraphs:
    for b in list(para._p.iter(qn('w:pBdr'))):b.getparent().remove(b)
    if para.style.name in ['Title','Subtitle']:
        for run in para.runs:
            run.font.name='Times New Roman'
            run.font.color.rgb=RGBColor(0,0,0)
            run.font.italic=False
for t in doc.tables:
    if len(t.rows)<=9:
        for row in t.rows[:-1]:
            for cell in row.cells:
                for para in cell.paragraphs:para.paragraph_format.keep_with_next=True
for idx,para in enumerate(doc.paragraphs):
    if para.text=='6.6 Subject Index':para.paragraph_format.page_break_before=True
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
dest=OUT/'Enterprise RAG Assistant SRS.docx'
doc.save(dest)
(ROOT/'tmp'/'srs-headings.json').write_text(json.dumps(headings,indent=2),encoding='utf-8')
print(dest)
print('Paragraphs',len(doc.paragraphs),'Tables',len(doc.tables),'FR records',len(reqs))
